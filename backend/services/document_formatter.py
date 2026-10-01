from __future__ import annotations

import html
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    """
    Normalize typographic characters that can cause
    formatting/export problems.
    """

    replacements = {
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "–": "-",
        "—": "-",
        "…": "...",
        "•": "-",
        "\u00a0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def split_sections(
    text: str,
) -> list[tuple[str, str]]:
    """
    Detect uppercase headings and split the document
    into heading/body sections.
    """

    clean_text = sanitize_text(text)

    lines = [
        line.rstrip()
        for line in clean_text.splitlines()
    ]

    sections: list[tuple[str, str]] = []

    current_heading = ""
    current_body: list[str] = []

    for line in lines:
        stripped = line.strip()

        is_heading = bool(stripped) and (
            (
                stripped.isupper()
                and len(stripped) <= 120
            )
            or re.match(
                r"^(SECTION|ARTICLE)\s+[A-Z0-9IVX.-]+",
                stripped,
                re.IGNORECASE,
            )
        )

        if is_heading:

            if current_heading or current_body:
                sections.append(
                    (
                        current_heading,
                        "\n".join(
                            current_body
                        ).strip(),
                    )
                )

            current_heading = stripped
            current_body = []

        else:
            current_body.append(line)

    if current_heading or current_body:
        sections.append(
            (
                current_heading,
                "\n".join(
                    current_body
                ).strip(),
            )
        )

    return [
        (heading, body)
        for heading, body in sections
        if heading or body
    ]


def _add_page_number(paragraph):
    """
    Add a Word PAGE field to the footer.
    """

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run("Page ")

    field = OxmlElement("w:fldSimple")
    field.set(
        qn("w:instr"),
        "PAGE",
    )

    run._r.addnext(field)


def format_docx(
    text: str,
    doc_type: str,
    logo_path: Path | None = None,
    footer_text: str = "",
) -> bytes:

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # Logo
    if (
        logo_path
        and logo_path.exists()
    ):
        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            str(logo_path),
            width=Inches(1.35),
        )

    # Main title
    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        sanitize_text(
            doc_type
        ).upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    # Content
    for heading, body in split_sections(text):

        if heading:

            paragraph = doc.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(4)

            run = paragraph.add_run(
                heading
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        for line in body.splitlines():

            if not line.strip():
                continue

            paragraph = doc.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.15

            run = paragraph.add_run(
                line.strip()
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # Footer
    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        footer_text
        or (
            "LegalEase - AI-assisted draft - "
            "review before use"
        )
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    _add_page_number(footer)

    output = BytesIO()

    doc.save(output)

    return output.getvalue()


class BrandedPDF(FPDF):

    def __init__(
        self,
        logo_path: Path | None,
        footer_text: str,
    ):

        super().__init__(
            orientation="P",
            unit="mm",
            format="A4",
        )

        self.logo_path = logo_path
        self.footer_text = footer_text

        self.set_auto_page_break(
            auto=True,
            margin=18,
        )

    def header(self):

        if (
            self.logo_path
            and self.logo_path.exists()
        ):
            self.image(
                str(self.logo_path),
                x=88,
                y=8,
                w=34,
            )

            self.ln(22)

        else:
            self.ln(8)

    def footer(self):

        self.set_y(-14)

        self.set_font(
            "Times",
            size=7,
        )

        self.cell(
            0,
            5,
            self.footer_text[:150],
            align="C",
        )

        self.set_y(-9)

        self.cell(
            0,
            5,
            f"Page {self.page_no()}",
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_path: Path | None = None,
    footer_text: str = "",
) -> bytes:

    pdf = BrandedPDF(
        logo_path,
        footer_text
        or (
            "LegalEase - AI-assisted draft - "
            "review before use"
        ),
    )

    pdf.set_title(
        sanitize_text(doc_type)
    )

    pdf.add_page()

    # Title
    pdf.set_font(
        "Times",
        "B",
        16,
    )

    pdf.multi_cell(
        0,
        9,
        sanitize_text(
            doc_type
        ).upper(),
        align="C",
    )

    pdf.ln(4)

    # Sections
    for heading, body in split_sections(text):

        if heading:

            pdf.set_font(
                "Times",
                "B",
                11,
            )

            pdf.multi_cell(
                0,
                6,
                heading,
            )

            pdf.ln(1)

        if body:

            pdf.set_font(
                "Times",
                size=10.5,
            )

            for line in body.splitlines():

                line = line.strip()

                if not line:
                    continue

                pdf.multi_cell(
                    0,
                    5.4,
                    line,
                )

                pdf.ln(0.8)

            pdf.ln(1.5)

    return bytes(
        pdf.output()
    )


def format_txt(
    text: str,
) -> bytes:

    return sanitize_text(
        text
    ).encode("utf-8")


def format_html_preview(
    text: str,
) -> str:

    escaped = html.escape(
        sanitize_text(text)
    )

    output = []

    for line in escaped.splitlines():

        if not line.strip():
            continue

        clean_line = line.strip()

        if (
            clean_line.isupper()
            and len(clean_line) <= 120
        ):
            output.append(
                f"<h3>{clean_line}</h3>"
            )

        else:
            output.append(
                f"<p>{line}</p>"
            )

    return "".join(output)